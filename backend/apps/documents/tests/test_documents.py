"""D4 (PDF de la liquidación) y D5 (Excel), con el almacenamiento simulado por moto."""
from io import BytesIO

import boto3
import pytest
from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.test import override_settings
from moto import mock_aws
from openpyxl import load_workbook
from pypdf import PdfReader

from apps.accounts.models import User
from apps.api.testing import api_client_for
from apps.documents.models import SettlementDocument
from apps.people.models import Occupancy, Person
from tests.factories import OrgMembershipFactory


pytestmark = pytest.mark.django_db

CONCEPTS = ('O.S.E. (agua)', 'U.T.E. (luz)', 'Limpieza', 'Honorarios de administración',
            'Fondo de reserva', 'Cuota extraordinaria')


@pytest.fixture
def s3():
    with override_settings(
        S3_ENDPOINT_URL='', S3_PUBLIC_ENDPOINT_URL='', S3_ACCESS_KEY_ID='test',
        S3_SECRET_ACCESS_KEY='test', S3_BUCKET='documentos-test',
    ), mock_aws():
        client = boto3.client('s3', region_name=settings.S3_REGION)
        client.create_bucket(Bucket='documentos-test')
        yield client


@pytest.fixture
def demo(monkeypatch, s3):
    monkeypatch.setenv('DEMO_PASSWORD', 'demo-password-123')
    with override_settings(DEBUG=True):
        call_command('seed_demo')
    return User.objects.get(email='maria@admin.demo')


def minas(user):
    return user.org_memberships.get().organization.buildings.get(name='Edificio Minas')


def occupancy_101():
    return Occupancy.objects.get(unit__building__name='Edificio Minas', unit__number='101', role='owner')


def workbook(response):
    assert response.status_code == 200
    assert response['Content-Type'].startswith('application/vnd.openxmlformats')
    return load_workbook(BytesIO(response.content))


def pdf_text(s3, document):
    body = s3.get_object(Bucket='documentos-test', Key=document.file_key)['Body'].read()
    return '\n'.join(page.extract_text() for page in PdfReader(BytesIO(body)).pages)


# ── D5: Excel ───────────────────────────────────────────────────────────
def test_accounts_export_has_one_row_per_account_with_numeric_amounts(demo):
    response = api_client_for(demo).get(f'/api/v1/buildings/{minas(demo).pk}/exports/accounts?month=2026-10')
    sheet = workbook(response).active

    assert 'cobros-edificio-minas-2026-10.xlsx' in response['Content-Disposition']
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 18
    juan = next(r for r in rows if r[0] == '101' and r[3] == 'Propietario')
    assert juan[8] == pytest.approx(4153.80)
    assert isinstance(juan[8], float)
    assert sheet.cell(row=2, column=9).number_format == '"$" #,##0.00'
    assert juan[2] == pytest.approx(0.086)


def test_text_that_looks_like_a_formula_is_written_as_text(demo):
    Person.objects.filter(pk=occupancy_101().person_id).update(
        first_name='=HYPERLINK("http://malo")', last_name='x'
    )
    response = api_client_for(demo).get(f'/api/v1/buildings/{minas(demo).pk}/exports/accounts?month=2026-10')
    rows = list(workbook(response).active.iter_rows(min_row=2, values_only=True))

    names = [r[4] for r in rows if r[0] == '101']
    assert names[0].startswith("'=HYPERLINK")


def test_exports_are_scoped(demo):
    outsider = api_client_for(OrgMembershipFactory().user)
    assert outsider.get(f'/api/v1/buildings/{minas(demo).pk}/exports/accounts').status_code == 404
    resident = api_client_for(User.objects.get(email='juan.perez@correo.com.uy'))
    assert resident.get(f'/api/v1/buildings/{minas(demo).pk}/exports/accounts').status_code == 403


def test_movements_export_for_staff_and_owner_only(demo):
    occupancy = occupancy_101()
    staff = workbook(api_client_for(demo).get(f'/api/v1/occupancies/{occupancy.pk}/exports/movements'))
    rows = list(staff.active.iter_rows(min_row=2, values_only=True))
    assert rows[-1][4] == pytest.approx(4153.80)

    juan = api_client_for(User.objects.get(email='juan.perez@correo.com.uy'))
    assert juan.get(f'/api/v1/portal/occupancies/{occupancy.pk}/exports/movements').status_code == 200
    other = api_client_for(User.objects.get(email='maria.garcia@correo.com.uy'))
    assert other.get(f'/api/v1/portal/occupancies/{occupancy.pk}/exports/movements').status_code == 404


def test_settlement_export_lists_concepts_and_prorrateo(demo):
    book = workbook(api_client_for(demo).get(f'/api/v1/buildings/{minas(demo).pk}/settlements/2026-10/export'))
    concepts = list(book['Conceptos'].iter_rows(min_row=2, values_only=True))
    prorrateo = list(book['Prorrateo'].iter_rows(min_row=2, values_only=True))

    assert [r[0] for r in concepts[:-1]] == list(CONCEPTS)
    assert concepts[-1][3] == pytest.approx(48300)
    assert sum(r[-1] for r in prorrateo) == pytest.approx(48300)
    assert len(prorrateo) == 18


# ── D4: PDF ─────────────────────────────────────────────────────────────
def test_unit_pdf_has_total_and_concepts(demo, s3):
    juan = api_client_for(User.objects.get(email='juan.perez@correo.com.uy'))
    response = juan.get(f'/api/v1/portal/occupancies/{occupancy_101().pk}/settlements/2026-10/pdf')

    assert response.status_code == 200
    assert 'X-Amz-Expires=300' in response.json()['url']
    text = pdf_text(s3, SettlementDocument.objects.get(occupancy=occupancy_101()))
    assert '$ 4.153,80' in text
    for concept in CONCEPTS:
        assert concept in text
    assert 'Apto 101' in text


def test_pdf_is_generated_once_and_never_replaced(demo, s3):
    client = api_client_for(demo)
    url = f'/api/v1/buildings/{minas(demo).pk}/settlements/2026-10/pdf'
    assert client.get(url).status_code == 200
    first = SettlementDocument.objects.get(occupancy=None)
    assert client.get(url).status_code == 200

    assert SettlementDocument.objects.filter(occupancy=None).count() == 1
    assert SettlementDocument.objects.get(occupancy=None).file_key == first.file_key
    assert s3.list_objects_v2(Bucket='documentos-test')['KeyCount'] == 1
    with pytest.raises(ValidationError):
        first.save()


def test_building_pdf_has_totals_but_no_resident_names(demo, s3):
    juan = api_client_for(User.objects.get(email='juan.perez@correo.com.uy'))
    response = juan.get(f'/api/v1/portal/occupancies/{occupancy_101().pk}/settlements/2026-10/building-pdf')

    assert response.status_code == 200
    text = pdf_text(s3, SettlementDocument.objects.get(occupancy=None))
    assert '$ 48.300,00' in text
    assert 'Juan Pérez' not in text
    assert 'María García' not in text


def test_pdfs_are_scoped(demo):
    occupancy = occupancy_101()
    outsider = api_client_for(OrgMembershipFactory().user)
    assert outsider.get(f'/api/v1/buildings/{minas(demo).pk}/settlements/2026-10/pdf').status_code == 404
    assert outsider.get(f'/api/v1/occupancies/{occupancy.pk}/settlements/2026-10/pdf').status_code == 404
    other = api_client_for(User.objects.get(email='maria.garcia@correo.com.uy'))
    assert other.get(f'/api/v1/portal/occupancies/{occupancy.pk}/settlements/2026-10/pdf').status_code == 404
    # Un mes sin liquidación emitida tampoco existe.
    assert api_client_for(demo).get(f'/api/v1/buildings/{minas(demo).pk}/settlements/2026-11/pdf').status_code == 404


def test_unit_pdf_previous_balance_is_taken_at_issue_date(demo, s3):
    # La 101 pagó setiembre el 8 de octubre, antes de que se emitiera octubre en la demo.
    juan = api_client_for(User.objects.get(email='juan.perez@correo.com.uy'))
    juan.get(f'/api/v1/portal/occupancies/{occupancy_101().pk}/settlements/2026-10/pdf')
    text = pdf_text(s3, SettlementDocument.objects.get(occupancy=occupancy_101()))

    assert 'Saldo anterior $ 0,00' in text
    assert 'Total a pagar antes del día 10 $ 4.153,80' in text
    assert 'maria@admin.demo' not in text
