"""Exportaciones a Excel (D5): contingencia analógica que pide el cliente.

Los montos llegan como Decimal (calculados por los mismos servicios que la API) y se
convierten a número solo al escribir la celda, con formato de moneda para que Excel sume.
"""
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MONEY_FORMAT = '"$" #,##0.00'
COEFFICIENT_FORMAT = '0.000000'
# Tope para exportar en el request; edificios más grandes irán por el worker (D2).
MAX_ROWS = 5000
FORMULA_START = ('=', '+', '-', '@', '\t', '\r')


def safe_text(value):
    """Evita inyección de fórmulas: un texto que Excel leería como fórmula va con apóstrofo."""
    text = '' if value is None else str(value)
    return "'" + text if text.startswith(FORMULA_START) else text


class Sheet:
    """Hoja con encabezado en negrita y columnas tipadas (texto, moneda, coeficiente, entero)."""

    def __init__(self, workbook, title, columns, *, first=False):
        self.ws = workbook.active if first else workbook.create_sheet()
        self.ws.title = title[:31]
        self.columns = columns
        self.ws.append([label for label, _ in columns])
        for cell in self.ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical='center')
        self.ws.freeze_panes = 'A2'
        for index, (label, _) in enumerate(columns, start=1):
            self.ws.column_dimensions[self.ws.cell(row=1, column=index).column_letter].width = max(12, len(label) + 4)

    def append(self, values):
        if self.ws.max_row > MAX_ROWS:
            raise ValueError(f'La exportación supera el tope de {MAX_ROWS} filas.')
        row = self.ws.max_row + 1
        for index, ((_, kind), value) in enumerate(zip(self.columns, values), start=1):
            cell = self.ws.cell(row=row, column=index)
            if kind == 'money':
                cell.value = float(Decimal(value or 0))
                cell.number_format = MONEY_FORMAT
            elif kind == 'coefficient':
                cell.value = float(Decimal(value or 0))
                cell.number_format = COEFFICIENT_FORMAT
            elif kind == 'int':
                cell.value = value
            else:
                cell.value = safe_text(value)
                cell.data_type = 's'


def response(workbook, filename):
    buffer = BytesIO()
    workbook.save(buffer)
    result = HttpResponse(buffer.getvalue(), content_type=XLSX)
    result['Content-Disposition'] = f'attachment; filename="{filename}"'
    return result


def file_slug(name):
    return slugify(name) or 'edificio'


STATUS_LABEL = {
    'pagado': 'Al día', 'revision': 'En revisión', 'pendiente': 'Pendiente',
    'moroso': 'Moroso', 'sin-cargar': 'Sin liquidación',
}
ROLE_LABEL = {'owner': 'Propietario', 'tenant': 'Inquilino'}


def accounts_workbook(building, month, rows, floors):
    """Cobros del edificio: una fila por cuenta (unidad + rol), como la pestaña Cobros."""
    workbook = Workbook()
    sheet = Sheet(workbook, f'Cobros {month:%Y-%m}', [
        ('Unidad', 'text'), ('Piso', 'int'), ('Coeficiente', 'coefficient'), ('Rol', 'text'),
        ('Responsable', 'text'), ('Teléfono', 'text'), ('Email', 'text'), ('Cargo del mes', 'money'),
        ('Saldo', 'money'), ('Estado', 'text'), ('Meses vencidos', 'text'),
    ], first=True)
    for row in rows:
        sheet.append([
            row['unit'], floors.get(row['unit']), row['coefficient'], ROLE_LABEL[row['role']],
            row['person']['name'], row['person']['phone'], row['person']['email'],
            row['charge'], row['balance'], STATUS_LABEL[row['status']], ' '.join(row['overdue_months']),
        ])
    return response(workbook, f'cobros-{file_slug(building.name)}-{month:%Y-%m}.xlsx')


def movements_workbook(occupancy, movements):
    """Cuenta corriente de una ocupación, del movimiento más viejo al más nuevo."""
    workbook = Workbook()
    sheet = Sheet(workbook, 'Cuenta corriente', [
        ('Fecha', 'text'), ('Concepto', 'text'), ('Debe', 'money'), ('Haber', 'money'), ('Saldo', 'money'),
    ], first=True)
    for item in reversed(movements):
        sheet.append([item['date'].isoformat(), item['concept'], item['debit'], item['credit'], item['balance']])
    unit = occupancy.unit
    return response(
        workbook,
        f'cuenta-{file_slug(unit.building.name)}-{slugify(unit.number)}-{occupancy.role}.xlsx',
    )


def settlement_workbook(settlement, items, charges):
    """Liquidación del mes: los conceptos y el prorrateo por unidad (una fila por cuenta)."""
    workbook = Workbook()
    concepts = Sheet(workbook, 'Conceptos', [
        ('Concepto', 'text'), ('Tipo', 'text'), ('Detalle', 'text'), ('Monto', 'money'),
    ], first=True)
    for item in items:
        concepts.append([
            item.concept.name,
            'Solo propietario' if item.concept.kind == 'owner_only' else 'Ordinario',
            item.detail, item.amount,
        ])
    concepts.append(['Total', '', '', sum((item.amount for item in items), Decimal('0'))])

    names = [item.concept.name for item in items]
    prorrateo = Sheet(
        workbook, 'Prorrateo',
        [('Unidad', 'text'), ('Rol', 'text'), ('Responsable', 'text'), ('Coeficiente', 'coefficient')]
        + [(name, 'money') for name in names] + [('Total', 'money')],
    )
    for account in charges:
        prorrateo.append(
            [account['unit'], ROLE_LABEL[account['role']], account['person'], account['coefficient']]
            + [account['by_concept'].get(name, Decimal('0')) for name in names]
            + [account['total']]
        )
    building = settlement.building
    return response(
        workbook, f'liquidacion-{file_slug(building.name)}-{settlement.month:%Y-%m}.xlsx'
    )
